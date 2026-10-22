"""
로컬 MariaDB(meal_rf_finger) 전용 데이터 액세스.
pc_app 화면들이 기대하는 메서드 집합을 제공한다 (HTTP/FastAPI 없음).
"""

from __future__ import annotations

import base64
import hashlib
import json
import os
import secrets
from datetime import date, datetime, time, timedelta, timezone
from typing import Any, Dict, List, Optional, Tuple

from meal_db.connection import get_connection

try:
    import mariadb as _mariadb_driver
except ImportError:
    _mariadb_driver = None  # type: ignore
from meal_db.f22_sync import (
    remove_employee_from_f22_terminals as _f22_remove_employee,
    sync_employee_to_f22_terminals as _f22_sync_employee,
    sync_fingerprint_to_f22_terminals as _f22_sync,
)

# 관리자 최초·신규 등록 시 기본 비밀번호 (로그인 구현 시 검증에 사용)
DEFAULT_ADMIN_INITIAL_PASSWORD = "meal1234!"
_ADMIN_PBKDF2_ITERS = 310000


def _hash_password(plain: str) -> str:
    salt = secrets.token_hex(16)
    h = hashlib.pbkdf2_hmac(
        "sha256",
        plain.encode("utf-8"),
        bytes.fromhex(salt),
        _ADMIN_PBKDF2_ITERS,
    )
    return "pbkdf2_sha256$%d$%s$%s" % (_ADMIN_PBKDF2_ITERS, salt, h.hex())


def _verify_password(plain: str, stored: Optional[str]) -> bool:
    if not stored or not plain:
        return False
    parts = stored.split("$")
    if len(parts) != 4 or parts[0] != "pbkdf2_sha256":
        return False
    try:
        iters = int(parts[1])
        salt = bytes.fromhex(parts[2])
        want = parts[3]
        h = hashlib.pbkdf2_hmac("sha256", plain.encode("utf-8"), salt, iters)
        return h.hex() == want
    except Exception:
        return False


def _iso(dt: Any) -> str:
    if isinstance(dt, datetime):
        if dt.tzinfo is None:
            return dt.isoformat()
        return dt.astimezone(timezone.utc).replace(tzinfo=None).isoformat() + "Z"
    if isinstance(dt, str):
        return dt
    return str(dt)


def _parse_created_at(val: Any) -> datetime:
    if isinstance(val, datetime):
        return val.replace(tzinfo=None) if val.tzinfo else val
    s = str(val).replace("Z", "").strip()
    if "T" in s:
        return datetime.fromisoformat(s.split("+")[0].split(".")[0])
    return datetime.strptime(s[:19], "%Y-%m-%d %H:%M:%S")


def _local_today() -> date:
    """Windows/OS에 설정된 로컬 타임존 기준 오늘. DB가 같은 PC에서 SYSTEM 타임존이면 DATE(created_at)와 달력일이 맞기 쉽다."""
    return datetime.now().astimezone().date()


def _format_mysql_time(val: Any) -> str:
    """MariaDB TIME → UI용 HH:MM:SS (Connector는 time 또는 timedelta로 줄 수 있음)."""
    if val is None:
        return ""
    if isinstance(val, str):
        return val.strip()
    if isinstance(val, datetime):
        return val.strftime("%H:%M:%S")
    if isinstance(val, time):
        return val.strftime("%H:%M:%S")
    if isinstance(val, timedelta):
        sec = int(val.total_seconds()) % 86400
        if sec < 0:
            sec += 86400
        h, r = divmod(sec, 3600)
        m, s = divmod(r, 60)
        return f"{h:02d}:{m:02d}:{s:02d}"
    return str(val)


def _policy_row(row: Dict[str, Any]) -> Dict[str, Any]:
    d = dict(row)
    d["start_time"] = _format_mysql_time(d.get("start_time"))
    d["end_time"] = _format_mysql_time(d.get("end_time"))
    return d


def _db_time_as_time(val: Any) -> time:
    """MariaDB TIME / Connector 반환값 → time."""
    if val is None:
        return time(0, 0, 0)
    if isinstance(val, time):
        return val
    if isinstance(val, timedelta):
        sec = int(val.total_seconds()) % 86400
        if sec < 0:
            sec += 86400
        h, r = divmod(sec, 3600)
        m, s = divmod(r, 60)
        return time(h, m, s)
    if isinstance(val, datetime):
        return val.time()
    s = str(val).strip()
    if not s:
        return time(0, 0, 0)
    parts = s.split(":")
    try:
        h = int(parts[0]) if parts else 0
        mi = int(parts[1]) if len(parts) > 1 else 0
        se = int(float(parts[2])) if len(parts) > 2 else 0
        return time(h, mi, se)
    except (ValueError, TypeError):
        return time(0, 0, 0)


def _time_in_policy_span(t: time, start_t: time, end_t: time) -> bool:
    if start_t <= end_t:
        return start_t <= t <= end_t
    return t >= start_t or t <= end_t


def terminal_attendance_stable_key(rec: Dict[str, Any]) -> str:
    """라이브·백필(폴링) 동일 건이 같은 auth_uid가 되도록 하는 안정 키(접미 순번 제외).

    시각은 초 단위로만 포함(마이크로초 제거)해 라이브/풀·재실행 시 키가 일치하도록 한다.
    zk_uid(단말 내부 사용자 슬롯)는 라이브/풀에서 달라지는 경우가 많아 키에 넣지 않는다.
    """
    ts = rec.get("timestamp")
    if not isinstance(ts, datetime):
        return ""
    if ts.tzinfo is not None:
        ts = ts.astimezone().replace(tzinfo=None)
    else:
        ts = ts.replace(tzinfo=None)
    ts = ts.replace(microsecond=0)
    user_id = (rec.get("user_id") or "").strip()
    device = str(rec.get("device") or "")
    ts_key = ts.strftime("%Y-%m-%d %H:%M:%S")
    return "%s|%s|%s|%s|%s" % (
        device,
        user_id,
        ts_key,
        rec.get("punch", 0),
        rec.get("status", 0),
    )


def terminal_auth_failure_stable_key(rec: Dict[str, Any]) -> str:
    ts = rec.get("timestamp")
    if not isinstance(ts, datetime):
        return ""
    if ts.tzinfo is not None:
        ts = ts.astimezone().replace(tzinfo=None)
    else:
        ts = ts.replace(tzinfo=None)
    ts = ts.replace(microsecond=0)
    return "%s|%s|%s|%s|%s|%s" % (
        str(rec.get("device") or ""),
        (rec.get("card_no") or "").strip(),
        (rec.get("pin_raw") or "").strip(),
        ts.strftime("%Y-%m-%d %H:%M:%S"),
        rec.get("event_code", 0),
        rec.get("adms_index", 0),
    )


_AUTH_FAILURE_TABLE_SQL = """
CREATE TABLE IF NOT EXISTS auth_failure_logs (
  id BIGINT AUTO_INCREMENT PRIMARY KEY,
  created_at DATETIME NOT NULL,
  card_no VARCHAR(128) NULL,
  terminal_device VARCHAR(128) NULL,
  terminal_label VARCHAR(128) NULL,
  adms_sn VARCHAR(64) NULL,
  event_code INT NULL,
  event_label VARCHAR(128) NULL,
  verify_type VARCHAR(64) NULL,
  pin_raw VARCHAR(64) NULL,
  auth_uid VARCHAR(128) NOT NULL,
  recorded_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
  UNIQUE KEY uq_auth_failure_auth_uid (auth_uid),
  KEY idx_auth_failure_created (created_at),
  KEY idx_auth_failure_card (card_no, created_at)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
"""


def _ensure_auth_failure_logs_table(cur) -> None:
    cur.execute(_AUTH_FAILURE_TABLE_SQL)
def _terminal_user_id_lookup_candidates(raw: str) -> List[str]:
    """단말 user_id(사번 또는 카드번호) 후보. 앞자리 0이 빠진 10자리 카드 등 대응."""
    u = (raw or "").strip()
    if not u:
        return []
    out: List[str] = []
    seen: set = set()

    def add(x: str) -> None:
        x = (x or "").strip()
        if x and x not in seen:
            seen.add(x)
            out.append(x)

    add(u)
    if u.isdigit():
        try:
            pad = int(os.environ.get("MEAL_ZK_USER_ID_ZERO_PAD", "0") or "0")
        except ValueError:
            pad = 0
        if pad > 0 and len(u) < pad:
            add(u.zfill(pad))
        if len(u) > 1 and u.startswith("0"):
            add(u.lstrip("0") or "0")
    return out


def _env_truthy(key: str, default: str = "1") -> bool:
    v = (os.environ.get(key, default) or default).strip().lower()
    return v in ("1", "true", "yes", "y", "on")


def _find_employee_for_terminal_user_id(cur, user_id: str) -> Optional[Dict[str, Any]]:
    for cand in _terminal_user_id_lookup_candidates(user_id):
        cur.execute(
            """
            SELECT emp_no, company_code, name FROM employees
            WHERE status = 'ACTIVE'
              AND (emp_no = %s OR rf_card_uid = %s OR TRIM(rf_card_uid) = %s)
            ORDER BY emp_no LIMIT 1
            """,
            (cand, cand, cand),
        )
        row = cur.fetchone()
        if row:
            return dict(row)
    return None


class LocalMealClient:
    """MariaDB Connector/Python 기반."""

    _is_local = True

    def __init__(self):
        pass

    def close(self):
        pass

    def _default_company_code(self, cur) -> str:
        cur.execute("SELECT code FROM companies ORDER BY code LIMIT 1")
        row = cur.fetchone()
        return str(row["code"]) if row and row.get("code") else "HQ"

    @staticmethod
    def _max_suffix_int(prefix: str, codes: List[str]) -> int:
        """prefix + 숫자 형태 코드 중 최대 번호 (없으면 0)."""
        p = (prefix or "").strip().upper()
        mx = 0
        for raw in codes:
            c = str(raw or "").strip()
            if len(c) <= len(p) or c[: len(p)].upper() != p:
                continue
            tail = c[len(p) :]
            if tail.isdigit():
                mx = max(mx, int(tail))
        return mx

    def _next_company_code(self, cur) -> str:
        cur.execute("SELECT code FROM companies")
        rows = cur.fetchall() or []
        codes = [str(r["code"]) for r in (rows or []) if r and r.get("code") is not None]
        n = self._max_suffix_int("C", codes) + 1
        if n <= 9999:
            return "C%04d" % n
        return "C%d" % n

    def _next_department_code(self, cur) -> str:
        cur.execute("SELECT code FROM departments")
        rows = cur.fetchall() or []
        codes = [str(r["code"]) for r in (rows or []) if r and r.get("code") is not None]
        n = self._max_suffix_int("D", codes) + 1
        if n <= 9999:
            return "D%04d" % n
        return "D%d" % n

    # --- Companies ---
    def get_companies(self):
        with get_connection() as conn:
            cur = conn.cursor()
            cur.execute("SELECT code, name FROM companies ORDER BY code")
            return list(cur.fetchall())

    def create_company(self, code, name):
        try:
            name = (name or "").strip()
            if not name:
                return False, "회사명을 입력하세요."
            with get_connection() as conn:
                cur = conn.cursor()
                c = self._next_company_code(cur)
                cur.execute("INSERT INTO companies (code, name) VALUES (%s, %s)", (c, name))
                cur.execute("SELECT code, name FROM companies WHERE code = %s", (c,))
                row = cur.fetchone()
                return True, dict(row) if row else {}
        except Exception as e:
            return False, str(e)

    def update_company(self, old_code, code, name):
        try:
            old_code = (old_code or "").strip()
            code = (code or "").strip()
            name = (name or "").strip()
            if not old_code or not code or not name:
                return None
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    "UPDATE companies SET code=%s, name=%s WHERE code=%s",
                    (code, name, old_code),
                )
                cur.execute("SELECT code, name FROM companies WHERE code = %s", (code,))
                row = cur.fetchone()
                return dict(row) if row else None
        except Exception:
            return None

    def delete_company(self, code):
        try:
            code = (code or "").strip()
            if not code:
                return False
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute("DELETE FROM companies WHERE code = %s", (code,))
                return cur.rowcount > 0
        except Exception:
            return False

    # --- Departments ---
    def get_departments(self, company_code=None):
        with get_connection() as conn:
            cur = conn.cursor()
            if company_code:
                cur.execute(
                    "SELECT code, company_code, name FROM departments WHERE company_code = %s ORDER BY code",
                    (str(company_code).strip(),),
                )
            else:
                cur.execute("SELECT code, company_code, name FROM departments ORDER BY code")
            return list(cur.fetchall())

    def create_department(self, company_code, code, name):
        try:
            name = (name or "").strip()
            company_code = (company_code or "").strip()
            if not name:
                return False, "부서명을 입력하세요."
            if not company_code:
                return False, "회사코드가 없습니다."
            with get_connection() as conn:
                cur = conn.cursor()
                c = self._next_department_code(cur)
                cur.execute(
                    "INSERT INTO departments (company_code, code, name) VALUES (%s, %s, %s)",
                    (company_code, c, name),
                )
                cur.execute(
                    "SELECT code, company_code, name FROM departments WHERE code = %s",
                    (c,),
                )
                row = cur.fetchone()
                return True, dict(row) if row else {}
        except Exception as e:
            return False, str(e)

    def update_department(self, old_code, code, name):
        try:
            old_code = (old_code or "").strip()
            code = (code or "").strip()
            name = (name or "").strip()
            if not old_code or not code or not name:
                return None
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    "UPDATE departments SET code=%s, name=%s WHERE code=%s",
                    (code, name, old_code),
                )
                cur.execute(
                    "SELECT code, company_code, name FROM departments WHERE code = %s",
                    (code,),
                )
                row = cur.fetchone()
                return dict(row) if row else None
        except Exception:
            return None

    def delete_department(self, code):
        try:
            code = (code or "").strip()
            if not code:
                return False
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    "UPDATE employees SET department_code = NULL WHERE department_code = %s",
                    (code,),
                )
                cur.execute("DELETE FROM departments WHERE code = %s", (code,))
                return cur.rowcount > 0
        except Exception:
            return False

    # --- Employees ---
    def _row_to_employee(self, row: dict, dept_name: str) -> dict:
        if not row:
            return {}
        out = {
            "company_code": row["company_code"],
            "department_code": row.get("department_code") or None,
            "emp_no": row["emp_no"],
            "name": row["name"],
            "status": row["status"],
            "rf_card_uid": row.get("rf_card_uid") or "",
            "card_uid": row.get("rf_card_uid") or "",
        }
        if "fp_has0" in row:
            out["fingerprint_has_slot0"] = bool(int(row["fp_has0"] or 0))
        if "fp_has1" in row:
            out["fingerprint_has_slot1"] = bool(int(row["fp_has1"] or 0))
        if "fp_has2" in row:
            out["fingerprint_has_slot2"] = bool(int(row["fp_has2"] or 0))
        if "fp_has3" in row:
            out["fingerprint_has_slot3"] = bool(int(row["fp_has3"] or 0))
        # SELECT * 경로(create/update 후)에만 있음 — F22 동기화 시 지문 템플릿 전달용
        if "fingerprint_template" in row:
            out["fingerprint_template_b64"] = row.get("fingerprint_template") or None
        if "fingerprint_template_2" in row:
            out["fingerprint_template_b64_2"] = row.get("fingerprint_template_2") or None
        if "fingerprint_template_3" in row:
            out["fingerprint_template_b64_3"] = row.get("fingerprint_template_3") or None
        if "fingerprint_template_4" in row:
            out["fingerprint_template_b64_4"] = row.get("fingerprint_template_4") or None
        for i in range(1, 5):
            k = "fingerprint_slot_fid_%d" % i
            if k in row:
                v = row.get(k)
                if v is None:
                    out[k] = None
                else:
                    try:
                        out[k] = int(v)
                    except (TypeError, ValueError):
                        out[k] = None
        return out

    def get_employees(self, search="", status=None):
        with get_connection() as conn:
            cur = conn.cursor()
            sql = """
                SELECT e.company_code, e.department_code, e.emp_no, e.name, e.status,
                       e.rf_card_uid,
                       (e.fingerprint_template IS NOT NULL) AS fp_has0,
                       (e.fingerprint_template_2 IS NOT NULL) AS fp_has1,
                       (e.fingerprint_template_3 IS NOT NULL) AS fp_has2,
                       (e.fingerprint_template_4 IS NOT NULL) AS fp_has3,
                       e.fingerprint_slot_fid_1, e.fingerprint_slot_fid_2,
                       e.fingerprint_slot_fid_3, e.fingerprint_slot_fid_4
                FROM employees e
                WHERE 1=1
            """
            args: List[Any] = []
            if status:
                sql += " AND e.status = %s"
                args.append(status)
            if search:
                sql += " AND (e.name LIKE %s OR e.emp_no LIKE %s)"
                pat = "%" + search + "%"
                args.extend([pat, pat])
            sql += " ORDER BY e.emp_no"
            cur.execute(sql, args)
            rows = cur.fetchall()
            cur.execute("SELECT code, name FROM departments")
            dm = {r["code"]: r["name"] for r in cur.fetchall()}
            out = []
            for r in rows:
                dn = dm.get(r.get("department_code"), "")
                out.append(self._row_to_employee(r, dn))
            return out

    def find_employees_by_rf_card(self, card_no: str):
        """카드번호(RF UID)로 사원·부서를 찾는다. 재직·퇴사 모두 포함."""
        card = str(card_no or "").strip()
        if not card or card == "-":
            return True, []
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                sql = """
                    SELECT e.emp_no, e.name, e.rf_card_uid, IFNULL(d.name, '') AS dept_name
                    FROM employees e
                    LEFT JOIN departments d ON d.code = e.department_code
                    WHERE TRIM(IFNULL(e.rf_card_uid, '')) = %s
                """
                args: List[Any] = [card]
                if card.isdigit():
                    sql += """
                        OR (
                            TRIM(IFNULL(e.rf_card_uid, '')) REGEXP '^[0-9]+$'
                            AND CAST(TRIM(e.rf_card_uid) AS UNSIGNED) = %s
                        )
                    """
                    args.append(int(card))
                sql += " ORDER BY e.emp_no"
                cur.execute(sql, args)
                rows = []
                for r in cur.fetchall() or []:
                    rows.append(
                        {
                            "emp_no": str(r.get("emp_no") or ""),
                            "name": str(r.get("name") or ""),
                            "dept_name": str(r.get("dept_name") or ""),
                            "rf_card_uid": str(r.get("rf_card_uid") or ""),
                        }
                    )
                return True, rows
        except Exception as e:
            return False, str(e)

    def _rf_card_uid_duplicate_message(self, cur, rf_uid, exclude_emp_no=None):
        """같은 UID가 다른 사원에게 있으면 저장 거절 문구. 빈 값은 허용."""
        uid = str(rf_uid or "").strip()
        if not uid:
            return ""
        sql = """
            SELECT emp_no, name FROM employees
            WHERE TRIM(rf_card_uid) = %s
        """
        args = [uid]
        exclude = str(exclude_emp_no or "").strip()
        if exclude:
            sql += " AND emp_no <> %s"
            args.append(exclude)
        sql += " ORDER BY emp_no LIMIT 1"
        cur.execute(sql, args)
        row = cur.fetchone()
        if not row:
            return ""
        return "UID가 중복되었습니다. (사번 %s, %s)" % (
            str(row.get("emp_no") or ""),
            str(row.get("name") or ""),
        )

    def get_active_employees_for_f22_sync(self) -> List[dict]:
        """재직 사원 전체 — F22 일괄 동기화용 (카드·지문 템플릿 포함)."""
        with get_connection() as conn:
            cur = conn.cursor()
            cur.execute(
                """
                SELECT e.*
                FROM employees e
                WHERE e.status = 'ACTIVE'
                ORDER BY e.emp_no
                """
            )
            rows = cur.fetchall() or []
            return [self._row_to_employee(r, "") for r in rows]

    def bulk_resync_active_employees_to_f22(
        self,
        terminals_only: list | None = None,
        *,
        cancel_check=None,
        on_progress=None,
    ) -> dict:
        from meal_db.f22_bulk_sync import bulk_resync_active_employees

        employees = self.get_active_employees_for_f22_sync()
        return bulk_resync_active_employees(
            self.get_device_settings(),
            employees,
            terminals_only,
            cancel_check=cancel_check,
            on_progress=on_progress,
        )

    def create_employee(self, data: dict):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cid = (data.get("company_code") or "").strip() or self._default_company_code(cur)
                did = (data.get("department_code") or "").strip() or None
                ft = data.get("fingerprint_template_b64") or None
                ft2 = data.get("fingerprint_template_b64_2") or None
                ft3 = data.get("fingerprint_template_b64_3") or None
                ft4 = data.get("fingerprint_template_b64_4") or None
                sf1 = data.get("fingerprint_slot_fid_1")
                sf2 = data.get("fingerprint_slot_fid_2")
                sf3 = data.get("fingerprint_slot_fid_3")
                sf4 = data.get("fingerprint_slot_fid_4")
                eno = str(data.get("emp_no") or "").strip()
                if not eno:
                    return False, "사번을 입력하세요."
                rf_uid = (str(data.get("rf_card_uid") or "").strip()) or None
                dup = self._rf_card_uid_duplicate_message(cur, rf_uid)
                if dup:
                    return False, dup
                cur.execute(
                    """INSERT INTO employees (company_code, department_code, emp_no, name, status, rf_card_uid,
                       fingerprint_template, fingerprint_template_2, fingerprint_template_3, fingerprint_template_4,
                       fingerprint_slot_fid_1, fingerprint_slot_fid_2, fingerprint_slot_fid_3, fingerprint_slot_fid_4)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (
                        cid,
                        did,
                        eno,
                        data["name"],
                        data.get("status") or "ACTIVE",
                        rf_uid,
                        ft,
                        ft2,
                        ft3,
                        ft4,
                        sf1,
                        sf2,
                        sf3,
                        sf4,
                    ),
                )
                cur.execute("SELECT * FROM employees WHERE emp_no = %s", (eno,))
                row = cur.fetchone()
                return True, self._row_to_employee(row, "")
        except Exception as e:
            return False, str(e)

    def update_employee(self, emp_no, data: dict):
        key = str(emp_no or "").strip()
        if not key:
            return False, "사번을 입력하세요."
        try:
            data = dict(data or {})
            lookup = key
            if "emp_no" in data:
                lookup = str(data.get("emp_no") or "").strip()
                if not lookup:
                    return False, "사번을 입력하세요."
                data["emp_no"] = lookup
            with get_connection() as conn:
                cur = conn.cursor()
                if "rf_card_uid" in data:
                    rf_uid = (str(data.get("rf_card_uid") or "").strip()) or None
                    data["rf_card_uid"] = rf_uid
                    dup = self._rf_card_uid_duplicate_message(cur, rf_uid, exclude_emp_no=key)
                    if dup:
                        return False, dup
                fields = []
                args: List[Any] = []
                mapping = {
                    "company_code": "company_code",
                    "department_code": "department_code",
                    "emp_no": "emp_no",
                    "name": "name",
                    "status": "status",
                    "rf_card_uid": "rf_card_uid",
                }
                for k, col in mapping.items():
                    if k not in data:
                        continue
                    val = data[k]
                    if k in ("company_code", "department_code", "emp_no"):
                        val = (str(val).strip() if val is not None else "") or None
                    fields.append("%s = %%s" % col)
                    args.append(val)
                if "fingerprint_template_b64" in data:
                    fields.append("fingerprint_template = %s")
                    args.append(data.get("fingerprint_template_b64") or None)
                if "fingerprint_template_b64_2" in data:
                    fields.append("fingerprint_template_2 = %s")
                    args.append(data.get("fingerprint_template_b64_2") or None)
                if "fingerprint_template_b64_3" in data:
                    fields.append("fingerprint_template_3 = %s")
                    args.append(data.get("fingerprint_template_b64_3") or None)
                if "fingerprint_template_b64_4" in data:
                    fields.append("fingerprint_template_4 = %s")
                    args.append(data.get("fingerprint_template_b64_4") or None)
                if "fingerprint_slot_fid_1" in data:
                    fields.append("fingerprint_slot_fid_1 = %s")
                    args.append(data.get("fingerprint_slot_fid_1"))
                if "fingerprint_slot_fid_2" in data:
                    fields.append("fingerprint_slot_fid_2 = %s")
                    args.append(data.get("fingerprint_slot_fid_2"))
                if "fingerprint_slot_fid_3" in data:
                    fields.append("fingerprint_slot_fid_3 = %s")
                    args.append(data.get("fingerprint_slot_fid_3"))
                if "fingerprint_slot_fid_4" in data:
                    fields.append("fingerprint_slot_fid_4 = %s")
                    args.append(data.get("fingerprint_slot_fid_4"))
                if not fields:
                    cur.execute("SELECT * FROM employees WHERE emp_no = %s", (key,))
                    row = cur.fetchone()
                    return True, self._row_to_employee(row, "")
                args.append(key)
                cur.execute(
                    "UPDATE employees SET " + ", ".join(fields) + " WHERE emp_no = %s",
                    args,
                )
                cur.execute("SELECT * FROM employees WHERE emp_no = %s", (lookup,))
                row = cur.fetchone()
                return True, self._row_to_employee(row, "")
        except Exception as e:
            return False, str(e)

    def delete_employee(self, emp_no, permanent=False):
        """DB 반영 후 등록된 F22 단말에서 해당 사번 사용자를 삭제한다. 반환: {\"db_ok\", \"f22_report\"}."""
        try:
            emp_no = str(emp_no or "").strip()
            if not emp_no:
                return {"db_ok": False, "f22_report": []}
            db_ok = False
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute("SELECT emp_no FROM employees WHERE emp_no = %s", (emp_no,))
                row = cur.fetchone()
                if not row:
                    return {"db_ok": False, "f22_report": []}
                if permanent:
                    cur.execute("DELETE FROM meal_logs WHERE emp_no = %s", (emp_no,))
                    cur.execute("DELETE FROM employees WHERE emp_no = %s", (emp_no,))
                    db_ok = cur.rowcount > 0
                else:
                    cur.execute(
                        "UPDATE employees SET status = 'RESIGNED' WHERE emp_no = %s",
                        (emp_no,),
                    )
                    db_ok = cur.rowcount > 0
            f22_report: List[Tuple[str, bool, str]] = []
            if db_ok and emp_no:
                f22_report = _f22_remove_employee(self.get_device_settings(), emp_no)
            return {"db_ok": db_ok, "f22_report": f22_report}
        except Exception:
            return {"db_ok": False, "f22_report": []}

    def import_employees_excel(self, company_code, file_content, replace=False):
        """사원일괄등록 양식(xlsx). 열: 사번, 이름, 부서코드 또는 부서명, RF카드 UID.

        replace 가 거짓이면 없는 사번만 추가한다.
        참이면 해당 회사 사원·식사기록을 지운 뒤 엑셀 행만 넣는다. 넣을 행이 없으면 삭제하지 않는다.
        """
        try:
            import openpyxl
            from io import BytesIO
        except ImportError:
            return False, "엑셀 읽기를 위해 openpyxl이 필요합니다."
        if not file_content:
            return False, "파일이 비어 있습니다."
        try:
            wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True, read_only=True)
        except Exception as e:
            return False, "엑셀 파일을 열 수 없습니다. xlsx 양식을 사용하세요.\n%s" % e
        try:
            ws = wb[wb.sheetnames[0]]
            grid = [list(row) for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        if not grid:
            return False, "시트가 비어 있습니다."

        def _norm_header(value) -> str:
            return str(value or "").strip().lower().replace(" ", "")

        aliases = {
            "사번": "emp_no",
            "사원번호": "emp_no",
            "empno": "emp_no",
            "emp_no": "emp_no",
            "이름": "name",
            "성명": "name",
            "name": "name",
            "부서코드": "department_code",
            "departmentcode": "department_code",
            "부서": "department_name",
            "부서명": "department_name",
            "rf카드uid": "rf_card_uid",
            "rf카드": "rf_card_uid",
            "카드uid": "rf_card_uid",
            "카드번호": "rf_card_uid",
            "상태": "status",
            "status": "status",
        }

        header_idx = None
        colmap = {}
        for i, row in enumerate(grid[:15]):
            mapped = {}
            for c, cell in enumerate(row):
                key = aliases.get(_norm_header(cell))
                if key and key not in mapped:
                    mapped[key] = c
            if "emp_no" in mapped and "name" in mapped:
                header_idx = i
                colmap = mapped
                break
        if header_idx is None:
            return False, "첫 행 근처에 사번, 이름 열이 없습니다. 사원일괄등록 양식을 사용하세요."

        def _text(value) -> str:
            if value is None:
                return ""
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            if isinstance(value, int):
                return str(value)
            return str(value).strip()

        def _col(row, key) -> str:
            idx = colmap.get(key)
            if idx is None or idx >= len(row):
                return ""
            return _text(row[idx])

        inserted = 0
        skipped = 0
        errors = []
        seen = set()
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cid = (company_code or "").strip() or self._default_company_code(cur)
                cur.execute("SELECT code, name FROM departments")
                by_code = {}
                by_name = {}
                for d in cur.fetchall() or []:
                    code = str(d.get("code") or "").strip()
                    name = str(d.get("name") or "").strip()
                    if code:
                        by_code[code] = code
                    if name and name not in by_name:
                        by_name[name] = code
                deleted = 0
                logs_deleted = 0
                if replace:
                    existing = set()
                    used_uids = {}
                else:
                    cur.execute("SELECT emp_no FROM employees")
                    existing = {str(r.get("emp_no") or "").strip() for r in (cur.fetchall() or [])}
                    cur.execute(
                        """SELECT emp_no, rf_card_uid FROM employees
                           WHERE rf_card_uid IS NOT NULL AND TRIM(rf_card_uid) <> ''"""
                    )
                    used_uids = {}
                    for r in cur.fetchall() or []:
                        uid = str(r.get("rf_card_uid") or "").strip()
                        eno_have = str(r.get("emp_no") or "").strip()
                        if uid and uid not in used_uids:
                            used_uids[uid] = eno_have
                pending = []

                for offset, row in enumerate(grid[header_idx + 1 :], start=header_idx + 2):
                    if not row or all(cell is None or str(cell).strip() == "" for cell in row):
                        continue
                    emp_no = _col(row, "emp_no")
                    name = _col(row, "name")
                    if emp_no.startswith("예시"):
                        continue
                    if not emp_no and not name:
                        continue
                    if not emp_no or not name:
                        errors.append("%d행: 사번과 이름을 모두 입력하세요." % offset)
                        continue
                    dept_code = _col(row, "department_code")
                    dept_name = _col(row, "department_name")
                    if dept_code:
                        if dept_code not in by_code:
                            errors.append("%d행: 부서코드 %s 가 없습니다." % (offset, dept_code))
                            continue
                    elif dept_name:
                        dept_code = by_name.get(dept_name) or ""
                        if not dept_code:
                            errors.append("%d행: 부서명 %s 가 없습니다." % (offset, dept_name))
                            continue
                    else:
                        dept_code = None
                    if emp_no in seen or emp_no in existing:
                        skipped += 1
                        errors.append("%d행: 사번 %s 은 이미 있습니다." % (offset, emp_no))
                        continue
                    rf_uid = _col(row, "rf_card_uid") or None
                    if rf_uid:
                        owner = used_uids.get(rf_uid)
                        if owner:
                            errors.append(
                                "%d행: UID %s 이(가) 중복되었습니다. (사번 %s)"
                                % (offset, rf_uid, owner)
                            )
                            continue
                        used_uids[rf_uid] = emp_no
                    status_raw = _col(row, "status")
                    if status_raw in ("퇴사", "RESIGNED", "resigned"):
                        status = "RESIGNED"
                    else:
                        status = "ACTIVE"
                    pending.append((cid, dept_code, emp_no, name, status, rf_uid))
                    seen.add(emp_no)
                    existing.add(emp_no)

                if pending:
                    if replace:
                        cur.execute(
                            """DELETE ml FROM meal_logs ml
                               INNER JOIN employees e ON e.emp_no = ml.emp_no
                               WHERE e.company_code = %s""",
                            (cid,),
                        )
                        logs_deleted = cur.rowcount
                        cur.execute(
                            "DELETE FROM employees WHERE company_code = %s",
                            (cid,),
                        )
                        deleted = cur.rowcount
                    for item in pending:
                        cur.execute(
                            """INSERT INTO employees
                               (company_code, department_code, emp_no, name, status, rf_card_uid)
                               VALUES (%s,%s,%s,%s,%s,%s)""",
                            item,
                        )
                        inserted += 1
        except Exception as e:
            return False, "등록 중 오류: %s" % e

        if replace and inserted == 0:
            note = "등록할 사원이 없어 기존 데이터는 삭제하지 않았습니다."
            if errors:
                note += "\n\n" + "\n".join(errors[:12])
                if len(errors) > 12:
                    note += "\n외 %d건" % (len(errors) - 12)
            return False, note

        if replace:
            lines = [
                "기존 사원 %d명, 식사기록 %d건 삭제 후 등록 %d건, 건너뜀 %d건"
                % (deleted, logs_deleted, inserted, skipped)
            ]
        else:
            lines = ["추가 등록 %d건, 건너뜀 %d건" % (inserted, skipped)]
        if errors:
            lines.append("")
            lines.extend(errors[:12])
            if len(errors) > 12:
                lines.append("외 %d건" % (len(errors) - 12))
        if inserted == 0 and errors:
            return False, "\n".join(lines)
        return True, {"message": "\n".join(lines), "inserted": inserted, "skipped": skipped}

    def reset_device_auth(self, emp_id):
        return True, {}

    # --- Device settings ---
    def get_device_settings(self):
        with get_connection() as conn:
            cur = conn.cursor()
            cur.execute("SELECT payload_json FROM device_settings WHERE id = 1")
            row = cur.fetchone()
            if not row:
                return {}
            try:
                return json.loads(row["payload_json"] or "{}")
            except json.JSONDecodeError:
                return {}

    def put_device_settings(self, data: dict):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    "UPDATE device_settings SET payload_json = %s WHERE id = 1",
                    (json.dumps(data, ensure_ascii=False),),
                )
                return True, data
        except Exception as e:
            return False, str(e)

    def get_attlog_clear_settings(self) -> Dict[str, Any]:
        """단말 출입 로그 동기화·삭제 지연(식사 종료 후 N분). 항상 스케줄 실행, 분만 설정."""
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    "SELECT minutes_after_meal_end FROM attlog_clear_settings WHERE id = 1"
                )
                row = cur.fetchone()
                if not row:
                    return {"minutes_after_meal_end": 10}
                return {
                    "minutes_after_meal_end": max(
                        0, min(24 * 60, int(row.get("minutes_after_meal_end") or 10))
                    ),
                }
        except Exception:
            return {"minutes_after_meal_end": 10}

    def put_attlog_clear_settings(self, minutes_after: int) -> Tuple[bool, str]:
        try:
            m = max(0, min(24 * 60, int(minutes_after)))
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    """
                    INSERT INTO attlog_clear_settings (id, enabled, minutes_after_meal_end)
                    VALUES (1, 1, %s)
                    ON DUPLICATE KEY UPDATE minutes_after_meal_end = VALUES(minutes_after_meal_end),
                      enabled = 1
                    """,
                    (m,),
                )
            return True, ""
        except Exception as e:
            return False, str(e)

    def get_active_policies_for_attlog_schedule(self) -> List[Dict[str, Any]]:
        """스케줄용: 활성 정책의 id·end_time (DB 원시 타입)."""
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    """
                    SELECT id, end_time FROM policies
                    WHERE is_active = 1
                      AND meal_type IS NOT NULL AND TRIM(meal_type) <> ''
                    ORDER BY id
                    """
                )
                return [dict(r) for r in cur.fetchall()]
        except Exception:
            return []

    def sync_fingerprint_to_f22_terminals(
        self,
        emp_no: str,
        display_name: str,
        rf_card_uid: str,
        fingerprint_uid: str,
        template_b64: str,
        template_b64_2: str | None = None,
        template_b64_3: str | None = None,
        template_b64_4: str | None = None,
        slot_fid_1: int | None = None,
        slot_fid_2: int | None = None,
        slot_fid_3: int | None = None,
        slot_fid_4: int | None = None,
    ):
        return _f22_sync(
            self.get_device_settings(),
            emp_no,
            display_name,
            rf_card_uid,
            fingerprint_uid,
            template_b64,
            template_b64_2,
            template_b64_3,
            template_b64_4,
            slot_fid_1,
            slot_fid_2,
            slot_fid_3,
            slot_fid_4,
        )

    def sync_employee_to_f22_terminals(
        self,
        emp_no: str,
        display_name: str,
        rf_card_uid: str,
        fingerprint_uid: str = "",
        template_b64: str | None = None,
        template_b64_2: str | None = None,
        template_b64_3: str | None = None,
        template_b64_4: str | None = None,
        slot_fid_1: int | None = None,
        slot_fid_2: int | None = None,
        slot_fid_3: int | None = None,
        slot_fid_4: int | None = None,
    ):
        return _f22_sync_employee(
            self.get_device_settings(),
            emp_no,
            display_name,
            rf_card_uid,
            fingerprint_uid,
            template_b64,
            template_b64_2,
            template_b64_3,
            template_b64_4,
            slot_fid_1,
            slot_fid_2,
            slot_fid_3,
            slot_fid_4,
        )

    def _pick_policy_for_terminal_punch(self, cur, company_code: str, dt: datetime) -> Tuple[Optional[int], int]:
        cur.execute(
            """
            SELECT id, start_time, end_time, base_price FROM policies
            WHERE is_active = 1 AND (company_code = %s OR company_code IS NULL)
            ORDER BY (company_code IS NULL), id
            """,
            (company_code,),
        )
        rows = cur.fetchall() or []
        t = dt.time()
        for r in rows:
            st = _db_time_as_time(r["start_time"])
            et = _db_time_as_time(r["end_time"])
            if _time_in_policy_span(t, st, et):
                return int(r["id"]), int(r["base_price"] or 0)
        return None, 0

    def ingest_terminal_attendance_record(
        self, rec: Dict[str, Any], dedup_salt: Optional[str] = None
    ) -> Dict[str, Any]:
        """단말 출입 1건 → meal_logs (단말에 남은 과거 포함 전부, auth_uid 중복 시 스킵).

        미래 시각만 제외. 폴링·실시간·ADMS PUSH 공통.
        """
        today = _local_today()
        ts = rec.get("timestamp")
        if not isinstance(ts, datetime):
            return {"inserted": 0, "skipped": 0, "unknown": 0}
        if ts.tzinfo is not None:
            ts = ts.astimezone().replace(tzinfo=None)
        else:
            ts = ts.replace(tzinfo=None)
        ts = ts.replace(microsecond=0)
        rec_day = ts.date()
        if rec_day > today:
            return {"inserted": 0, "skipped": 0, "unknown": 0, "out_of_range": 1}
        user_id = (rec.get("user_id") or "").strip()
        if not user_id:
            return {"inserted": 0, "skipped": 0, "unknown": 1}
        with get_connection() as conn:
            cur = conn.cursor()
            er = _find_employee_for_terminal_user_id(cur, user_id)
            if not er:
                return {"inserted": 0, "skipped": 0, "unknown": 1}
            emp_no = str(er.get("emp_no") or "").strip()
            emp_name = str(er.get("name") or "")
            cid = str(er["company_code"] or "")
            key_src = terminal_attendance_stable_key(rec)
            if not key_src:
                return {"inserted": 0, "skipped": 0, "unknown": 0}
            if dedup_salt is not None and str(dedup_salt).strip() != "":
                key_src += "|%s" % str(dedup_salt).strip()
            auth_uid = "zk:" + hashlib.sha256(key_src.encode("utf-8")).hexdigest()[:40]
            cur.execute("SELECT id FROM meal_logs WHERE auth_uid = %s LIMIT 1", (auth_uid,))
            if cur.fetchone():
                return {"inserted": 0, "skipped": 1, "unknown": 0}
            pid, price = self._pick_policy_for_terminal_punch(cur, cid, ts)
            try:
                cur.execute(
                    """INSERT INTO meal_logs (emp_no, policy_id, created_at, guest_count, final_price, is_void, auth_uid, path)
                       VALUES (%s,%s,%s,0,%s,0,%s,'terminal')""",
                    (emp_no, pid, ts, price, auth_uid),
                )
            except Exception as e:
                if _mariadb_driver is not None and isinstance(
                    e, _mariadb_driver.IntegrityError
                ):
                    try:
                        conn.rollback()
                    except Exception:
                        pass
                    return {"inserted": 0, "skipped": 1, "unknown": 0}
                raise
            meal_type = ""
            if pid:
                cur.execute("SELECT meal_type FROM policies WHERE id = %s", (pid,))
                pr = cur.fetchone()
                if pr:
                    meal_type = str(pr.get("meal_type") or "")
            date_time_str = ts.strftime("%Y-%m-%d %H:%M:%S")
            return {
                "inserted": 1,
                "skipped": 0,
                "unknown": 0,
                "is_today": rec_day == today,
                "emp_no": emp_no,
                "emp_name": emp_name,
                "meal_type": meal_type,
                "date_time_str": date_time_str,
            }

    def ingest_terminal_auth_failure_record(self, rec: Dict[str, Any]) -> Dict[str, Any]:
        """RF 카드 인증 실패 rtlog 1건 → auth_failure_logs."""
        ts = rec.get("timestamp")
        if not isinstance(ts, datetime):
            return {"inserted": 0, "skipped": 0}
        if ts.tzinfo is not None:
            ts = ts.astimezone().replace(tzinfo=None)
        else:
            ts = ts.replace(tzinfo=None)
        ts = ts.replace(microsecond=0)
        if ts.date() > _local_today():
            return {"inserted": 0, "skipped": 0, "out_of_range": 1}
        key_src = terminal_auth_failure_stable_key(rec)
        if not key_src:
            return {"inserted": 0, "skipped": 0}
        auth_uid = "af:" + hashlib.sha256(key_src.encode("utf-8")).hexdigest()[:40]
        card_no = (rec.get("card_no") or "").strip() or None
        with get_connection() as conn:
            cur = conn.cursor()
            _ensure_auth_failure_logs_table(cur)
            cur.execute("SELECT id FROM auth_failure_logs WHERE auth_uid = %s LIMIT 1", (auth_uid,))
            if cur.fetchone():
                return {"inserted": 0, "skipped": 1}
            try:
                cur.execute(
                    """INSERT INTO auth_failure_logs
                       (created_at, card_no, terminal_device, terminal_label, adms_sn,
                        event_code, event_label, verify_type, pin_raw, auth_uid)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (
                        ts,
                        card_no,
                        str(rec.get("device") or "") or None,
                        str(rec.get("label") or "") or None,
                        str(rec.get("adms_sn") or "") or None,
                        rec.get("event_code"),
                        str(rec.get("event_label") or "") or None,
                        str(rec.get("verify_type") or "") or None,
                        (rec.get("pin_raw") or "").strip() or None,
                        auth_uid,
                    ),
                )
            except Exception as e:
                if _mariadb_driver is not None and isinstance(
                    e, _mariadb_driver.IntegrityError
                ):
                    try:
                        conn.rollback()
                    except Exception:
                        pass
                    return {"inserted": 0, "skipped": 1}
                raise
            return {
                "inserted": 1,
                "skipped": 0,
                "card_no": card_no or "",
                "event_label": str(rec.get("event_label") or ""),
                "date_time_str": ts.strftime("%Y-%m-%d %H:%M:%S"),
            }

    def delete_auth_failure_log(self, log_id):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute("DELETE FROM auth_failure_logs WHERE id = %s", (log_id,))
                return True, cur.rowcount > 0
        except Exception as e:
            return False, str(e)

    def get_auth_failure_logs(self, search="", start_date=None, end_date=None):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                _ensure_auth_failure_logs_table(cur)
                sql = """
                    SELECT id, created_at, card_no, terminal_device, terminal_label,
                           adms_sn, event_code, event_label, verify_type, pin_raw, auth_uid
                    FROM auth_failure_logs
                    WHERE DATE(created_at) BETWEEN %s AND %s
                """
                args: List[Any] = [start_date or "1970-01-01", end_date or "2099-12-31"]
                if search:
                    sql += " AND (card_no LIKE %s OR event_label LIKE %s OR terminal_label LIKE %s)"
                    pat = "%" + search + "%"
                    args.extend([pat, pat, pat])
                sql += " ORDER BY created_at DESC"
                cur.execute(sql, args)
                rows = cur.fetchall()
            out = []
            for r in rows:
                row = dict(r)
                ts = row.get("created_at")
                if isinstance(ts, datetime):
                    row["created_at"] = ts.strftime("%Y-%m-%dT%H:%M:%S")
                out.append(row)
            return (True, out)
        except Exception as e:
            return (False, str(e))

    def sync_attendance_from_f22_terminals(self) -> Dict[str, Any]:
        """F22 단말 attendance → meal_logs (단말 버퍼 전체, auth_uid 해시로 중복 방지)."""
        import f22_pyzk

        settings = self.get_device_settings()
        records, err_lines, pull_ok_devices = f22_pyzk.pull_attendance_logs(settings)
        today = _local_today()
        inserted = 0
        skipped = 0
        unknown = 0
        seen_total = 0
        stable_seen: Dict[str, int] = {}
        for rec in records:
            ts = rec.get("timestamp")
            if isinstance(ts, datetime):
                tsn = ts.astimezone().replace(tzinfo=None) if ts.tzinfo else ts.replace(tzinfo=None)
                if tsn.date() <= today:
                    seen_total += 1
            sk = terminal_attendance_stable_key(rec)
            if sk:
                n = stable_seen.get(sk, 0)
                stable_seen[sk] = n + 1
                suff: Optional[str] = None if n == 0 else str(n)
            else:
                suff = None
            r = self.ingest_terminal_attendance_record(rec, dedup_salt=suff)
            inserted += int(r.get("inserted") or 0)
            skipped += int(r.get("skipped") or 0)
            unknown += int(r.get("unknown") or 0)
        cleared_labels: List[str] = []
        if _env_truthy("MEAL_ZK_CLEAR_ATTLOG_AFTER_PULL", "1") and pull_ok_devices:
            try:
                pwd = int(os.environ.get("ZK_DEVICE_COMM_PASSWORD", "0") or "0")
            except ValueError:
                pwd = 0
            for dev in pull_ok_devices:
                ok, cmsg = f22_pyzk.clear_terminal_attlog(
                    str(dev.get("ip") or ""),
                    int(dev.get("port") or 4370),
                    pwd,
                )
                lab = str(dev.get("label") or dev.get("ip") or "")
                if ok:
                    cleared_labels.append(lab)
                else:
                    err_lines.append("[%s] 단말 출퇴근 로그 삭제 실패: %s" % (lab, cmsg))
        detail = (
            "단말 기록: %d건 조회, 신규 반영 %d건, 건너뜀(중복) %d건, 미등록 사번 %d건"
            % (seen_total, inserted, skipped, unknown)
        )
        if cleared_labels:
            detail += " | 단말 버퍼 삭제: %s" % ", ".join(cleared_labels)
        if err_lines:
            detail += " | " + "; ".join(err_lines[:3])
            if len(err_lines) > 3:
                detail += "…"
        return {
            "inserted": inserted,
            "skipped": skipped,
            "unknown_user": unknown,
            "seen_total": seen_total,
            "errors": err_lines,
            "detail": detail,
        }

    # --- Stats & raw data ---
    def get_stats(self):
        day = _local_today()
        day_str = day.isoformat()
        with get_connection() as conn:
            cur = conn.cursor()
            cur.execute(
                "SELECT COUNT(*) AS c FROM meal_logs WHERE DATE(created_at) = %s AND is_void = 0",
                (day_str,),
            )
            total = int(cur.fetchone()["c"])
            cur.execute(
                """
                SELECT COALESCE(p.meal_type, '번외') AS meal_type, COUNT(*) AS cnt
                FROM meal_logs m
                LEFT JOIN policies p ON m.policy_id = p.id
                WHERE DATE(m.created_at) = %s AND m.is_void = 0
                GROUP BY COALESCE(p.meal_type, '번외')
                """,
                (day_str,),
            )
            counts: Dict[str, int] = {}
            for r in cur.fetchall():
                mt = str(r["meal_type"] or "번외")
                counts[mt] = int(r["cnt"])
            # 활성 정책의 식사명은 오늘 로그 0건이어도 대시보드 카드에 표시
            cur.execute(
                """
                SELECT meal_type, MIN(id) AS sort_id
                FROM policies
                WHERE is_active = 1 AND meal_type IS NOT NULL AND TRIM(meal_type) <> ''
                GROUP BY meal_type
                ORDER BY sort_id
                """
            )
            ordered: List[str] = []
            seen_mt: set = set()
            for r in cur.fetchall():
                mt = str(r["meal_type"]).strip()
                if not mt or mt in seen_mt:
                    continue
                seen_mt.add(mt)
                ordered.append(mt)
            for mt in counts:
                if mt not in seen_mt:
                    seen_mt.add(mt)
                    ordered.append(mt)
            # policy_id 없음 → SQL에서 '번외'로 묶임. 대시보드는 exception_count 카드 한 곳만 쓰므로 여기서는 제외
            summaries = [
                {"meal_type": mt, "count": counts.get(mt, 0)}
                for mt in ordered
                if mt != "번외"
            ]
            cur.execute(
                """
                SELECT COUNT(*) AS c FROM meal_logs
                WHERE DATE(created_at) = %s AND is_void = 0 AND policy_id IS NULL
                """,
                (day_str,),
            )
            exc = int(cur.fetchone()["c"])
        return {
            "date": day_str,
            "total_count": total,
            "meal_summaries": summaries,
            "exception_count": exc,
        }

    def _log_row_to_api(self, row: dict) -> dict:
        user = {
            "name": row.get("emp_name") or "",
            "emp_no": str(row.get("emp_no") or ""),
            "department_code": row.get("department_code") or None,
            "department_name": row.get("dept_name") or "",
            "rf_card_uid": row.get("rf_card_uid") or "",
        }
        policy = {}
        if row.get("meal_type"):
            policy = {"meal_type": row["meal_type"]}
        ca = row.get("created_at")
        return {
            "id": row["id"],
            "created_at": _iso(ca) if ca else "",
            "user": user,
            "policy": policy,
            "policy_id": row.get("policy_id"),
            "emp_no": str(row.get("emp_no") or ""),
            "path": row.get("path") or "local",
            "is_void": bool(row.get("is_void")),
            "guest_count": int(row.get("guest_count") or 0),
            "final_price": int(row.get("final_price") or 0),
            "auth_uid": row.get("auth_uid") or "",
        }

    def get_raw_data(
        self,
        search="",
        start_date=None,
        end_date=None,
        *,
        non_void_only: bool = False,
    ):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                sql = """
                    SELECT m.id, m.emp_no, m.policy_id, m.created_at, m.guest_count, m.final_price,
                           m.is_void, m.void_reason, m.auth_uid, m.path,
                           e.name AS emp_name, e.rf_card_uid, e.department_code,
                           d.name AS dept_name, p.meal_type, p.base_price
                    FROM meal_logs m
                    LEFT JOIN employees e ON m.emp_no = e.emp_no
                    LEFT JOIN departments d ON e.department_code = d.code
                    LEFT JOIN policies p ON m.policy_id = p.id
                    WHERE DATE(m.created_at) BETWEEN %s AND %s
                """
                args: List[Any] = [start_date or "1970-01-01", end_date or "2099-12-31"]
                if non_void_only:
                    sql += " AND m.is_void = 0"
                if search:
                    sql += " AND (e.name LIKE %s OR e.emp_no LIKE %s)"
                    pat = "%" + search + "%"
                    args.extend([pat, pat])
                sql += " ORDER BY m.created_at DESC"
                cur.execute(sql, args)
                rows = cur.fetchall()
            return (True, [self._log_row_to_api(dict(r)) for r in rows])
        except Exception as e:
            return (False, str(e))

    def create_manual_raw_data(self, data):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                emp_no = str(data.get("emp_no") or "").strip()
                if not emp_no:
                    return False, "사번이 없습니다."
                pid = data.get("policy_id")
                guest = int(data.get("guest_count") or 0)
                created = _parse_created_at(data["created_at"])
                price = 0
                if pid:
                    cur.execute("SELECT base_price FROM policies WHERE id = %s", (pid,))
                    pr = cur.fetchone()
                    if pr:
                        price = int(pr["base_price"] or 0)
                cur.execute(
                    """INSERT INTO meal_logs (emp_no, policy_id, created_at, guest_count, final_price, is_void, path)
                       VALUES (%s,%s,%s,%s,%s,0,'manual')""",
                    (emp_no, pid, created, guest, price),
                )
                lid = cur.lastrowid
                cur.execute(
                    """
                    SELECT m.id, m.emp_no, m.policy_id, m.created_at, m.guest_count, m.final_price,
                           m.is_void, m.auth_uid, m.path,
                           e.name AS emp_name, e.rf_card_uid, e.department_code,
                           d.name AS dept_name, p.meal_type, p.base_price
                    FROM meal_logs m
                    JOIN employees e ON m.emp_no = e.emp_no
                    LEFT JOIN departments d ON e.department_code = d.code
                    LEFT JOIN policies p ON m.policy_id = p.id
                    WHERE m.id = %s
                    """,
                    (lid,),
                )
                return True, self._log_row_to_api(dict(cur.fetchone()))
        except Exception as e:
            return False, str(e)

    def update_raw_data(self, log_id, data):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                emp_no = str(data.get("emp_no") or "").strip()
                if not emp_no:
                    return False, "사번이 없습니다."
                pid = data.get("policy_id")
                guest = int(data.get("guest_count") or 0)
                created = _parse_created_at(data["created_at"])
                price = 0
                if pid:
                    cur.execute("SELECT base_price FROM policies WHERE id = %s", (pid,))
                    pr = cur.fetchone()
                    if pr:
                        price = int(pr["base_price"] or 0)
                cur.execute(
                    """UPDATE meal_logs SET emp_no=%s, policy_id=%s, created_at=%s, guest_count=%s, final_price=%s
                       WHERE id = %s""",
                    (emp_no, pid, created, guest, price, log_id),
                )
                cur.execute(
                    """
                    SELECT m.id, m.emp_no, m.policy_id, m.created_at, m.guest_count, m.final_price,
                           m.is_void, m.auth_uid, m.path,
                           e.name AS emp_name, e.rf_card_uid, e.department_code,
                           d.name AS dept_name, p.meal_type, p.base_price
                    FROM meal_logs m
                    JOIN employees e ON m.emp_no = e.emp_no
                    LEFT JOIN departments d ON e.department_code = d.code
                    LEFT JOIN policies p ON m.policy_id = p.id
                    WHERE m.id = %s
                    """,
                    (log_id,),
                )
                return True, self._log_row_to_api(dict(cur.fetchone()))
        except Exception as e:
            return False, str(e)

    def delete_raw_data(self, log_id):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute("DELETE FROM meal_logs WHERE id = %s", (log_id,))
                return cur.rowcount > 0
        except Exception:
            return False

    def void_log(self, log_id, reason):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    "UPDATE meal_logs SET is_void = 1, void_reason = %s WHERE id = %s",
                    (reason or "", log_id),
                )
                return cur.rowcount > 0
        except Exception:
            return False

    # --- Policies ---
    def get_policies(self):
        with get_connection() as conn:
            cur = conn.cursor()
            cur.execute(
                """SELECT id, company_code, meal_type, start_time, end_time,
                          base_price, guest_price, is_active
                   FROM policies ORDER BY id"""
            )
            return [_policy_row(dict(r)) for r in cur.fetchall()]

    def create_policy(self, data: dict):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cid = (data.get("company_code") or "").strip() or None
                if cid is None:
                    cid = self._default_company_code(cur)
                cur.execute(
                    """INSERT INTO policies (company_code, meal_type, start_time, end_time, base_price, guest_price, is_active)
                       VALUES (%s,%s,%s,%s,%s,%s,%s)""",
                    (
                        cid,
                        data["meal_type"],
                        data["start_time"],
                        data["end_time"],
                        int(data.get("base_price") or 0),
                        int(data.get("guest_price") or data.get("base_price") or 0),
                        1 if data.get("is_active", True) else 0,
                    ),
                )
                pid = cur.lastrowid
                cur.execute(
                    """SELECT id, company_code, meal_type, start_time, end_time,
                              base_price, guest_price, is_active
                       FROM policies WHERE id = %s""",
                    (pid,),
                )
                row = cur.fetchone()
                return True, _policy_row(dict(row)) if row else {}
        except Exception as e:
            return False, str(e)

    def update_policy(self, policy_id, data: dict):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    """UPDATE policies SET meal_type=%s, start_time=%s, end_time=%s, base_price=%s, guest_price=%s, is_active=%s
                       WHERE id = %s""",
                    (
                        data["meal_type"],
                        data["start_time"],
                        data["end_time"],
                        int(data.get("base_price") or 0),
                        int(data.get("guest_price") or data.get("base_price") or 0),
                        1 if data.get("is_active", True) else 0,
                        policy_id,
                    ),
                )
                cur.execute(
                    """SELECT id, company_code, meal_type, start_time, end_time,
                              base_price, guest_price, is_active
                       FROM policies WHERE id = %s""",
                    (policy_id,),
                )
                row = cur.fetchone()
                if not row:
                    return False, "정책을 찾을 수 없습니다."
                return True, _policy_row(dict(row))
        except Exception as e:
            return False, str(e)

    def delete_policy(self, policy_id):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute("UPDATE meal_logs SET policy_id = NULL WHERE policy_id = %s", (policy_id,))
                cur.execute("DELETE FROM policies WHERE id = %s", (policy_id,))
                return cur.rowcount > 0
        except Exception:
            return False

    # --- Admins ---
    def get_admins(self):
        with get_connection() as conn:
            cur = conn.cursor()
            cur.execute("SELECT id, emp_no, name, is_verified FROM admins ORDER BY id")
            rows = cur.fetchall()
            return [
                {"id": r["id"], "emp_no": r["emp_no"], "name": r["name"], "is_verified": bool(r["is_verified"])}
                for r in rows
            ]

    def create_admin(self, emp_no: str, name: str):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                ph = _hash_password(DEFAULT_ADMIN_INITIAL_PASSWORD)
                cur.execute(
                    "INSERT INTO admins (emp_no, name, password_hash, is_verified) VALUES (%s, %s, %s, 0)",
                    (emp_no, name, ph),
                )
                aid = cur.lastrowid
                cur.execute("SELECT id, emp_no, name, is_verified FROM admins WHERE id = %s", (aid,))
                r = cur.fetchone()
                return True, {"id": r["id"], "emp_no": r["emp_no"], "name": r["name"], "is_verified": bool(r["is_verified"])}
        except Exception as e:
            return False, str(e)

    def create_admin_auto_login_id(self, name: str):
        """로그인 ID(emp_no) 자동 부여: admin → admin2 → admin3 …"""
        nm = (name or "").strip()
        if not nm:
            return False, "이름을 입력하세요."
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute("SELECT emp_no FROM admins")
                used = {str(r["emp_no"]).strip() for r in (cur.fetchall() or []) if r.get("emp_no")}
                if "admin" not in used:
                    emp_no = "admin"
                else:
                    emp_no = ""
                    for n in range(2, 100000):
                        cand = "admin%d" % n
                        if cand not in used:
                            emp_no = cand
                            break
                    if not emp_no:
                        return False, "로그인 ID를 더 이상 자동 생성할 수 없습니다."
                ph = _hash_password(DEFAULT_ADMIN_INITIAL_PASSWORD)
                cur.execute(
                    "INSERT INTO admins (emp_no, name, password_hash, is_verified) VALUES (%s, %s, %s, 0)",
                    (emp_no, nm, ph),
                )
                aid = cur.lastrowid
                cur.execute("SELECT id, emp_no, name, is_verified FROM admins WHERE id = %s", (aid,))
                r = cur.fetchone()
                return True, {"id": r["id"], "emp_no": r["emp_no"], "name": r["name"], "is_verified": bool(r["is_verified"])}
        except Exception as e:
            return False, str(e)

    def update_admin(self, user_id: int, name: str):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute("UPDATE admins SET name = %s WHERE id = %s", (name, user_id))
                return (cur.rowcount > 0, None)
        except Exception as e:
            return (False, str(e))

    def delete_admin(self, user_id: int):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute("DELETE FROM admins WHERE id = %s", (user_id,))
                return cur.rowcount > 0
        except Exception:
            return False

    def reset_admin_device(self, user_id: int):
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute("UPDATE admins SET is_verified = 0 WHERE id = %s", (user_id,))
                return cur.rowcount >= 0
        except Exception:
            return False

    def verify_admin_password(self, emp_no: str, password: str) -> bool:
        """사번(로그인 ID) + 비밀번호 검증 (향후 앱 로그인용)."""
        u = (emp_no or "").strip()
        if not u:
            return False
        try:
            with get_connection() as conn:
                cur = conn.cursor()
                cur.execute(
                    "SELECT password_hash FROM admins WHERE emp_no = %s LIMIT 1",
                    (u,),
                )
                row = cur.fetchone()
                if not row:
                    return False
                return _verify_password(password or "", row.get("password_hash"))
        except Exception:
            return False

    def get_excel_report_data(self, year, month):
        return None
